from datetime import date
import csv
import io
import re

from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.db.models import Q
from django.contrib.auth.decorators import login_required
from django.utils import timezone

from .models import Tenant
from .forms import TenantForm
from invoices.models import Invoice
from payments.models import Payment
from water_bills.models import WaterBill, WaterBillPayment
from properties.models import Property
from units.models import Unit
from leases.models import Lease
from PROPATIA.pagination import paginate_queryset
from accounts.access_utils import get_accessible_properties, filter_tenants_by_accessible_properties


def is_app_admin(user):
    return user.is_authenticated and hasattr(user, 'profile') and user.profile.role == 'admin'


@login_required
def tenant_list(request):
    if request.method == "POST":
        form = TenantForm(request.POST, request.FILES, user=request.user)
        if form.is_valid():
            tenant = form.save(commit=False)
            unit = form.cleaned_data.get('unit')
            if unit:
                tenant.unit = unit
                tenant.property = unit.property
                unit.status = 'occupied'
                unit.save()
                tenant.status = 'active'
            else:
                tenant.status = 'inactive'
            tenant.save()
            if tenant.unit:
                Lease.objects.create(
                    user=request.user,
                    tenant=tenant,
                    unit=tenant.unit,
                    start_date=timezone.now().date(),
                    monthly_rent=tenant.unit.rent_amount,
                    deposit_held=0,
                    is_active=True,
                )
            return redirect('tenants:tenant_list')
    else:
        form = TenantForm(user=request.user)

    accessible_props = get_accessible_properties(request.user)
    tenants = Tenant.objects.filter(
        Q(unit__property__in=accessible_props) | Q(property__in=accessible_props) | Q(unit__isnull=True)
    ).select_related('unit', 'unit__property', 'property').order_by('last_name', 'first_name')

    property_filter = request.GET.get('property')
    status_filter = request.GET.get('status')
    allowed_property_ids = set(accessible_props.values_list('id', flat=True))

    if property_filter and property_filter.isdigit() and int(property_filter) in allowed_property_ids:
        prop_id = int(property_filter)
        tenants = tenants.filter(
            Q(unit__property_id=prop_id) | Q(property_id=prop_id) | Q(unit__isnull=True)
        )
    else:
        property_filter = None

    if status_filter:
        tenants = tenants.filter(status=status_filter)

    pagination = paginate_queryset(request, tenants)

    return render(request, 'tenants/tenants_view.html', {
        'tenants': pagination['page_obj'],
        'form': form,
        'properties': accessible_props,
        'selected_property': property_filter,
        'selected_status': status_filter,
        'can_edit_tenants': is_app_admin(request.user),
        **pagination,
    })


@login_required
def available_units(request, pk):
    """Return units for a property as JSON."""
    accessible_props = get_accessible_properties(request.user)
    property_obj = get_object_or_404(accessible_props, pk=pk)
    units = Unit.objects.filter(property=property_obj).values('id', 'name', 'rent_amount')
    return JsonResponse({'property': property_obj.name, 'units': list(units)})


@login_required
def tenant_ledger(request, pk):
    """Render the tenant ledger with invoices and payments for the tenant/unit."""
    accessible_props = get_accessible_properties(request.user)
    tenant = get_object_or_404(
        Tenant.objects.select_related('unit', 'unit__property', 'property'),
        Q(pk=pk) & (Q(unit__property__in=accessible_props) | Q(property__in=accessible_props) | Q(unit__isnull=True)),
    )

    invoices = Invoice.objects.filter(tenant=tenant).select_related('unit').order_by('due_date', 'id')
    payments = Payment.objects.filter(tenant=tenant).select_related('unit').order_by('date', 'id')
    tenant_property = tenant.get_display_property()
    property_tenants = Tenant.objects.filter(
        Q(unit__property=tenant_property) | Q(property=tenant_property),
    ).select_related('unit', 'property').order_by('unit__name', 'last_name', 'first_name')

    ledger_entries = []
    for invoice in invoices:
        ledger_entries.append({
            'date': invoice.due_date,
            'code': '',
            'description': f"Invoice {invoice.invoice_number}",
            'debit': invoice.amount,
            'credit': 0,
            'status': invoice.status,
            'is_overdue': invoice.due_date < date.today() and invoice.status != 'Paid',
            'type': 'invoice',
        })

    for payment in payments:
        ledger_entries.append({
            'date': payment.date,
            'code': payment.code or '',
            'description': payment.description or 'Payment received',
            'debit': 0,
            'credit': payment.amount,
            'status': payment.status,
            'is_overdue': False,
            'type': 'payment',
        })

    water_bills = WaterBill.objects.filter(tenant=tenant).select_related('unit').order_by('due_date', 'id')
    for water_bill in water_bills:
        ledger_entries.append({
            'date': water_bill.due_date,
            'code': '',
            'description': f"Water Bill {water_bill.id}",
            'debit': water_bill.amount,
            'credit': 0,
            'status': water_bill.status,
            'is_overdue': water_bill.due_date < date.today() and water_bill.status != 'Paid',
            'type': 'water_bill',
        })

    water_bill_payments = WaterBillPayment.objects.filter(tenant=tenant).select_related('unit').order_by('date', 'id')
    for wb_payment in water_bill_payments:
        ledger_entries.append({
            'date': wb_payment.date,
            'code': wb_payment.code or '',
            'description': wb_payment.description or 'Water bill payment',
            'debit': 0,
            'credit': wb_payment.amount,
            'status': wb_payment.status,
            'is_overdue': False,
            'type': 'water_payment',
        })

    ledger_entries.sort(key=lambda item: (item['date'], item['type'] not in ('invoice', 'water_bill')))

    running_balance = 0
    for entry in ledger_entries:
        running_balance += entry['debit'] - entry['credit']
        entry['balance'] = running_balance

    total_invoiced = sum(invoice.amount for invoice in invoices) + sum(wb.amount for wb in water_bills)
    total_paid = sum(payment.amount for payment in payments) + sum(wbp.amount for wbp in water_bill_payments)
    current_balance = total_invoiced - total_paid
    overdue_count = sum(
        1 for invoice in invoices if invoice.due_date < date.today() and invoice.status != 'Paid'
    ) + sum(
        1 for wb in water_bills if wb.due_date < date.today() and wb.status != 'Paid'
    )

    return render(request, 'tenants/tenant_ledger.html', {
        'tenant': tenant,
        'ledger_entries': ledger_entries,
        'total_invoiced': total_invoiced,
        'total_paid': total_paid,
        'current_balance': current_balance,
        'overdue_count': overdue_count,
        'property_tenants': property_tenants,
    })


@login_required
def delete_tenants(request):
    """Delete selected tenants."""
    if request.method == 'POST':
        accessible_props = get_accessible_properties(request.user)
        tenant_ids = request.POST.getlist('tenant_ids[]')
        deleted_count, _ = Tenant.objects.filter(
            id__in=tenant_ids,
        ).filter(
            Q(unit__property__in=accessible_props) | Q(property__in=accessible_props) | Q(unit__isnull=True)
        ).delete()
        return JsonResponse({'success': True, 'message': f'{deleted_count} tenant(s) deleted'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})


@login_required
def edit_tenant(request, pk):
    accessible_props = get_accessible_properties(request.user)
    tenant = get_object_or_404(
        Tenant.objects.select_related('unit', 'unit__property', 'property'),
        Q(pk=pk) & (Q(unit__property__in=accessible_props) | Q(property__in=accessible_props) | Q(unit__isnull=True)),
    )

    if request.method == 'POST':
        if not is_app_admin(request.user):
            return JsonResponse({'success': False, 'message': 'Only admins can edit tenants'}, status=403)

        form = TenantForm(request.POST, request.FILES, instance=tenant, user=request.user)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True, 'message': 'Tenant updated'})
        return JsonResponse({'success': False, 'errors': form.errors}, status=400)

    return JsonResponse({
        'success': True,
        'tenant_id': tenant.id,
        'first_name': tenant.first_name,
        'last_name': tenant.last_name,
        'phone_number': tenant.phone_number,
        'next_of_kin_name': tenant.next_of_kin_name or '',
        'next_of_kin_phone_number': tenant.next_of_kin_phone_number or '',
        'description': tenant.description or '',
        'deposit_required': tenant.deposit_required,
        'deposit_amount': float(tenant.deposit_amount) if tenant.deposit_amount else 0,
        'property_id': tenant.property_id or (tenant.unit.property_id if tenant.unit else ''),
        'kra_pin': tenant.kra_pin or '',
        'id_card_front_url': tenant.id_card_front.url if tenant.id_card_front else '',
        'id_card_back_url': tenant.id_card_back.url if tenant.id_card_back else '',
    })


@login_required
def upload_tenants(request):
    """Upload tenants from CSV or XLSX file"""
    def normalize_header(header):
        return ' '.join(
            str(header)
            .replace('\ufeff', '')
            .replace('*', '')
            .replace('\\', '')
            .lower()
            .replace('_', ' ')
            .split()
        )

    def get_row_value(row, *headers):
        for header in headers:
            value = row.get(normalize_header(header))
            if value is not None and str(value).strip() != '':
                return value
        return None

    def normalize_upload_row(headers, values):
        row = {}
        for idx, header in enumerate(headers):
            normalized_header = normalize_header(header)
            if not normalized_header:
                continue
            value = values[idx] if idx < len(values) else None
            current_value = row.get(normalized_header)
            if current_value is None or str(current_value).strip() == '':
                row[normalized_header] = value
        return row

    if 'file' not in request.FILES:
        return JsonResponse({'success': False, 'message': 'No file provided'})

    file = request.FILES['file']
    filename = file.name.lower()
    validate_only = str(request.POST.get('validate_only', '')).lower() in {'1', 'true', 'yes', 'on'}

    try:
        rows = []

        if filename.endswith('.csv'):
            content = file.read().decode('utf-8-sig')
            sample = content[:2048]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',\t;')
            except csv.Error:
                dialect = csv.excel_tab if '\t' in sample else csv.excel
            reader = csv.reader(content.splitlines(), dialect=dialect)
            headers = next(reader, [])
            for line in reader:
                if any(str(v).strip() for v in line):
                    rows.append(normalize_upload_row(headers, line))
        elif filename.endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
            except ImportError:
                return JsonResponse({'success': False, 'message': 'openpyxl is not installed. Please use CSV format or contact admin.'})
            wb = load_workbook(io.BytesIO(file.read()), data_only=True)
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(v is not None and str(v).strip() != '' for v in row):
                    rows.append(normalize_upload_row(headers, list(row)))
        else:
            return JsonResponse({'success': False, 'message': 'Invalid file format. Please use CSV or XLSX'})

        if not rows:
            return JsonResponse({'success': False, 'message': 'File is empty'})

        accessible_props = get_accessible_properties(request.user)
        created_count = 0
        valid_rows = 0
        errors = []

        for idx, row in enumerate(rows, start=2):
            try:
                first_name = get_row_value(row, 'first name', 'firstname', 'first')
                last_name = get_row_value(row, 'last name', 'lastname', 'last')
                phone_number = get_row_value(row, 'phone number', 'phone', 'phone_number')
                property_name = get_row_value(row, 'property')
                unit_name = get_row_value(row, 'unit')
                next_of_kin_name = get_row_value(row, 'next of kin name', 'next_of_kin', 'kin name')
                next_of_kin_phone = get_row_value(row, 'next of kin phone', 'next_of_kin_phone_number', 'kin phone')
                description = get_row_value(row, 'description')
                deposit_required = get_row_value(row, 'deposit required', 'deposit_required')
                deposit_amount = get_row_value(row, 'deposit amount', 'deposit_amount')
                kra_pin = get_row_value(row, 'kra pin', 'kra_pin')

                # Validate KRA PIN format if provided (same regex as TenantForm)
                if kra_pin is not None:
                    kra_pin_val = str(kra_pin).strip()
                    if kra_pin_val and not re.match(r'^[A-Za-z]\d{9}[A-Za-z]$', kra_pin_val):
                        errors.append(f'Row {idx}: Invalid KRA PIN format (expected A123456789B)')
                        continue
                    kra_pin = kra_pin_val.upper() if kra_pin_val else None
                else:
                    kra_pin = None

                if not first_name or not last_name or not phone_number or not property_name:
                    errors.append(f'Row {idx}: Missing required fields (first_name, last_name, phone_number, property)')
                    continue

                property_obj = accessible_props.filter(name__iexact=str(property_name).strip()).first()
                if not property_obj:
                    errors.append(f'Row {idx}: Property "{property_name}" not found or not accessible')
                    continue

                unit = None
                if unit_name:
                    unit = Unit.objects.filter(property=property_obj, name__iexact=str(unit_name).strip()).first()
                    if not unit:
                        errors.append(f'Row {idx}: Unit "{unit_name}" not found in property "{property_name}"')
                        continue
                    if unit.status != 'vacant':
                        errors.append(f'Row {idx}: Unit "{unit_name}" in property "{property_name}" is not vacant')
                        continue
                else:
                    unit = Unit.objects.filter(property=property_obj, status='vacant').first()

                deposit_required_val = str(deposit_required).strip().lower() if deposit_required else 'false'
                deposit_required_bool = deposit_required_val in ('true', '1', 'yes', 'on')
                deposit_amount_val = 0
                if deposit_amount:
                    try:
                        deposit_amount_val = float(str(deposit_amount).replace(',', '').strip())
                    except (ValueError, TypeError):
                        deposit_amount_val = 0

                tenant_data = {
                    'property': property_obj,
                    'first_name': str(first_name).strip(),
                    'last_name': str(last_name).strip(),
                    'phone_number': str(phone_number).strip(),
                    'next_of_kin_name': str(next_of_kin_name).strip() if next_of_kin_name else '',
                    'next_of_kin_phone_number': str(next_of_kin_phone).strip() if next_of_kin_phone else '',
                    'description': str(description).strip() if description else '',
                    'deposit_required': deposit_required_bool,
                    'deposit_amount': deposit_amount_val,
                    'kra_pin': kra_pin,
                }

                if not unit:
                    # No vacant unit found — store property directly on the tenant
                    tenant_data['unit'] = None
                    tenant_data['status'] = 'inactive'
                    if validate_only:
                        valid_rows += 1
                        continue
                    tenant = Tenant.objects.create(**tenant_data)
                    created_count += 1
                    errors.append(f'Row {idx}: No vacant unit available in property "{property_name}" — tenant created without a unit')
                    continue

                # Unit found — assign it
                tenant_data['unit'] = unit
                tenant_data['status'] = 'active'
                if validate_only:
                    valid_rows += 1
                    continue

                tenant = Tenant.objects.create(**tenant_data)
                unit.status = 'occupied'
                unit.save()
                Lease.objects.create(
                    user=request.user,
                    tenant=tenant,
                    unit=unit,
                    start_date=timezone.now().date(),
                    monthly_rent=unit.rent_amount,
                    deposit_held=0,
                    is_active=True,
                )
                created_count += 1

            except ValueError as ve:
                errors.append(f'Row {idx}: Invalid data format - {str(ve)}')
            except Exception as e:
                errors.append(f'Row {idx}: {str(e)}')

        if validate_only:
            message = f'Validation complete: {valid_rows} valid, {len(errors)} invalid.'
            if errors:
                message += ' ' + '; '.join(errors[:3])
            response = {
                'success': not errors,
                'count': 0,
                'message': message,
                'errors': errors,
                'valid_rows': valid_rows,
                'invalid_rows': len(errors),
            }
            return JsonResponse(response)

        message = f'Successfully processed {created_count} tenants'
        if errors:
            message += f'. {len(errors)} errors: ' + '; '.join(errors[:3])

        response = {
            'success': created_count > 0 or not errors,
            'count': created_count,
            'message': message,
            'errors': errors,
            'invalid_rows': len(errors),
            'valid_rows': created_count,
        }
        return JsonResponse(response)

    except Exception as e:
        return JsonResponse({'success': False, 'message': f'Error processing file: {str(e)}'})


@login_required
def attach_tenant(request, pk):
    """Attach a unit to a tenant."""
    accessible_props = get_accessible_properties(request.user)
    tenant = get_object_or_404(Tenant, pk=pk)

    if request.method == 'POST':
        unit_id = request.POST.get('unit_id')
        if not unit_id:
            return JsonResponse({'success': False, 'message': 'No unit selected'}, status=400)

        unit = get_object_or_404(Unit, pk=unit_id)

        if unit.property not in accessible_props:
            return JsonResponse({'success': False, 'message': 'Unit not in an accessible property'}, status=403)

        if unit.status != 'vacant':
            return JsonResponse({'success': False, 'message': f'Unit {unit.name} is not vacant'}, status=400)

        if Lease.objects.filter(unit=unit, is_active=True).exists():
            return JsonResponse({'success': False, 'message': f'Unit {unit.name} already has an active lease'}, status=400)

        if tenant.leases.filter(is_active=True).exists():
            return JsonResponse({'success': False, 'message': 'Tenant already has an active lease'}, status=400)

        Lease.objects.create(
            user=request.user,
            tenant=tenant,
            unit=unit,
            start_date=timezone.now().date(),
            monthly_rent=unit.rent_amount,
            deposit_held=0,
            is_active=True,
        )
        tenant.unit = unit
        tenant.property = unit.property
        tenant.status = 'active'
        tenant.save()

        return JsonResponse({'success': True, 'message': f'Tenant attached to {unit.name}'})

    return JsonResponse({'success': False, 'message': 'Invalid request'}, status=400)


@login_required
def detach_tenant(request, pk):
    """Detach a tenant from their unit."""
    accessible_props = get_accessible_properties(request.user)
    tenant = get_object_or_404(
        Tenant.objects.select_related('unit', 'unit__property', 'property'),
        Q(pk=pk) & (Q(unit__property__in=accessible_props) | Q(property__in=accessible_props) | Q(unit__isnull=True))
    )

    if request.method == 'POST':
        lease = Lease.objects.filter(unit=tenant.unit, is_active=True).first()
        if lease:
            lease.is_active = False
            lease.save()

        if tenant.unit:
            tenant.unit.status = 'vacant'
            tenant.unit.save()

        tenant.unit = None
        tenant.status = 'inactive'
        tenant.save()

        return JsonResponse({'success': True, 'message': 'Tenant detached from unit'})

    return JsonResponse({'success': False, 'message': 'Invalid request'}, status=400)


@login_required
def attach_available_units(request):
    """Return all vacant units across accessible properties as JSON."""
    accessible_props = get_accessible_properties(request.user)
    units = Unit.objects.filter(
        property__in=accessible_props,
        status='vacant',
    ).values('id', 'name', 'rent_amount', 'property__name')

    return JsonResponse({'units': list(units)})
